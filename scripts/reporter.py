#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""报告输出：Markdown / Excel / JSON"""
import datetime as dt
import json
import os


def _fmt(v, n=1, dash="-"):
    if v is None:
        return dash
    try:
        return f"{float(v):.{n}f}"
    except (TypeError, ValueError):
        return dash


# ------------------------------------------------------------------ JSON
def write_json_result(path, ctx):
    """供分时采集器读取。只输出通过硬门槛的标的。"""
    passed = ctx["passed"]
    payload = {
        "status": "ok" if ctx["valued"] else "empty",
        "run_date": ctx["date"],
        "generated_at": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "model": ctx["cfg"]["model_version"],
        "data_valid": True,
        "candidate_pool_count": ctx["pool"],
        "passed_safety_margin_count": len(passed),
        "stocks": [
            {
                "code": x["code"],
                "name": x["name"],
                "price": x["price"],
                "score": x["score"],
                "dcf_value": x["dcf_value"],
                "dcf_discount": x["dcf_discount"],
                "graham_value": x["graham_value"],
                "graham_discount": x["graham_discount"],
                "pe_ttm": x["pe_ttm"],
                "pb": x["pb"],
                "roe_3y": x["roe_3y"],
                "np_cagr": x["np_cagr"],
                "market_cap": x["market_cap"],
            }
            for x in passed
        ],
        "pipeline": {
            "universe": ctx["universe"],
            "prefiltered": ctx["prefiltered"],
            "targets": ctx["targets"],
            "pool": ctx["pool"],
            "valued": len(ctx["valued"]),
        },
        "failed_count": len(ctx["stats"]["failed_stocks"]),
    }
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


# ------------------------------------------------------------------ Markdown
def write_markdown(path, ctx):
    cfg = ctx["cfg"]
    passed = ctx["passed"]
    valued = ctx["valued"]
    s = cfg["thresholds"]

    L = []
    A = L.append
    A(f"# 每周选股筛选 — 6指标估值模型 {cfg['model_version']}")
    A("")
    A(f"> 运行日期：**{ctx['date']}**　｜　模型口径见 `config/model_v3.md`　｜　"
      f"耗时 {ctx['elapsed']:.0f}s")
    A("")
    A("## 一、结果概览")
    A("")
    A("| 环节 | 数量 |")
    A("|---|---|")
    A(f"| 全A股 | {ctx['universe']} |")
    A(f"| 预筛（市值>100亿 + PE>0） | {ctx['prefiltered']} |")
    A(f"| 财务精算 | {ctx['targets']} |")
    A(f"| **粗筛池（九条件通过）** | **{ctx['pool']}** |")
    A(f"| 完成估值 | {len(valued)} |")
    A(f"| **通过安全边际（硬门槛）** | **{len(passed)}** |")
    A("")
    A(f"硬门槛：DCF 折价率 ≥ {s['min_dcf_discount']}% 且 现价 < 格林厄姆内在价值。")
    A("")

    # Top 5
    A("## 二、前 5 名")
    A("")
    if valued:
        A("| 排名 | 代码 | 名称 | 总分 | DCF折价率 | 现价 | DCF价值 | 格林厄姆 | 硬门槛 |")
        A("|---|---|---|---|---|---|---|---|---|")
        for i, x in enumerate(valued[:5], 1):
            mark = "✅ 通过" if x.get("pass") else "❌ 未过"
            A(f"| {i} | {x['code']} | {x['name']} | **{x['score']:.1f}** | "
              f"{_fmt(x['dcf_discount'])}% | {x['price']:.2f} | {_fmt(x['dcf_value'],2)} | "
              f"{_fmt(x['graham_value'],2)} | {mark} |")
    else:
        A("_无估值结果_")
    A("")

    # 通过列表
    A("## 三、通过安全边际的标的")
    A("")
    if passed:
        A("| # | 代码 | 名称 | 总分 | DCF折价率 | 格林厄姆折价率 | 现价 | ROE3Y | 净利CAGR | PE | 市值(亿) |")
        A("|---|---|---|---|---|---|---|---|---|---|---|")
        for i, x in enumerate(passed, 1):
            A(f"| {i} | {x['code']} | {x['name']} | {x['score']:.1f} | "
              f"**{_fmt(x['dcf_discount'])}%** | {_fmt(x['graham_discount'])}% | "
              f"{x['price']:.2f} | {_fmt(x['roe_3y'])} | {_fmt(x['np_cagr'])}% | "
              f"{_fmt(x['pe_ttm'])} | {_fmt(x['market_cap'],0)} |")
    else:
        A("> **本周期无标的通过硬门槛。**")
        A(">")
        A(f"> 粗筛池 {ctx['pool']} 只完成估值后，没有股票同时满足「DCF 折价 ≥20%」与「现价 < 格林厄姆价值」。")
        A("> 这本身是一个有价值的信号：当前池内标的整体估值不便宜，模型建议**空仓等待**，而非降低标准凑数。")
    A("")

    # 全量
    A("## 四、候选池全量排名")
    A("")
    if valued:
        A("| # | 代码 | 名称 | 总分 | DCF折价% | 格氏折价% | 现价 | DCF值 | 格氏值 | ROE3Y | CAGR% | PE | PB | 负债率% | 门槛 |")
        A("|---|---|---|---|---|---|---|---|---|---|---|---|---|---|---|")
        for i, x in enumerate(valued, 1):
            A(f"| {i} | {x['code']} | {x['name']} | {x['score']:.1f} | "
              f"{_fmt(x['dcf_discount'])} | {_fmt(x['graham_discount'])} | {x['price']:.2f} | "
              f"{_fmt(x['dcf_value'],2)} | {_fmt(x['graham_value'],2)} | {_fmt(x['roe_3y'])} | "
              f"{_fmt(x['np_cagr'])} | {_fmt(x['pe_ttm'])} | {_fmt(x['pb'],2)} | "
              f"{_fmt(x['debt_ratio'])} | {'✅' if x.get('pass') else '❌'} |")
    A("")

    # 未通过硬门槛明细
    near = [x for x in valued if not x.get("pass")]
    if near:
        A("### 接近门槛的标的（DCF折价 10%~20%）")
        A("")
        cand = [x for x in near if x["dcf_discount"] is not None and 10 <= x["dcf_discount"] < 20]
        if cand:
            A("| 代码 | 名称 | DCF折价% | 未通过原因 |")
            A("|---|---|---|---|")
            for x in cand[:10]:
                A(f"| {x['code']} | {x['name']} | {_fmt(x['dcf_discount'])} | {x['reason']} |")
        else:
            A("_无_")
        A("")

    # 失败项
    A("## 五、失败项")
    A("")
    fails = ctx["stats"]["failed_stocks"]
    A(f"**采集/估值失败：{len(fails)} 只**")
    A("")
    if fails:
        A("| 代码 | 名称 | 原因 |")
        A("|---|---|---|")
        for f in fails[:40]:
            A(f"| {f['code']} | {f['name']} | {f['reason']} |")
    else:
        A("_无_")
    A("")
    A(f"**粗筛剔除：{len(ctx['rejected'])} 只**（明细见附录）")
    A("")

    # 附录
    A("## 附录：粗筛剔除明细（前 60）")
    A("")
    A("| 代码 | 名称 | 未通过条件 |")
    A("|---|---|---|")
    for r in ctx["rejected"][:60]:
        A(f"| {r['code']} | {r['name']} | {r['reason']} |")
    A("")
    A("---")
    A("")
    A("> 本报告由模型自动生成，为估值筛选工具输出，不构成投资建议。"
      "DCF 以归母净利润替代自由现金流，对重资产行业存在高估倾向，详见模型文档第七章。")

    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(L))


# ------------------------------------------------------------------ Excel
def write_xlsx(path, ctx):
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return False

    cfg = ctx["cfg"]
    valued = ctx["valued"]
    passed = ctx["passed"]

    wb = openpyxl.Workbook()
    HDR = PatternFill("solid", fgColor="1F4E79")
    HDR_F = Font(color="FFFFFF", bold=True, size=10)
    PASS_FILL = PatternFill("solid", fgColor="E2EFDA")

    # --- Sheet1 概览
    ws = wb.active
    ws.title = "概览"
    ws["A1"] = f"6指标估值模型 {cfg['model_version']} — 选股筛选"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws.merge_cells("A1:D1")

    rows = [
        ("运行日期", ctx["date"]),
        ("全A股", ctx["universe"]),
        ("预筛(市值>100亿+PE>0)", ctx["prefiltered"]),
        ("财务精算", ctx["targets"]),
        ("粗筛池(九条件)", ctx["pool"]),
        ("完成估值", len(valued)),
        ("通过安全边际", len(passed)),
        ("耗时(秒)", round(ctx["elapsed"], 1)),
        ("", ""),
        ("硬门槛", f"DCF折价≥{cfg['thresholds']['min_dcf_discount']}% 且 现价<格林厄姆价值"),
    ]
    r = 3
    for k, v in rows:
        ws.cell(r, 1, k).font = Font(bold=True)
        ws.cell(r, 2, v)
        r += 1
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50

    # --- Sheet2 通过
    def sheet(name, items, cols, title):
        s = wb.create_sheet(name)
        s["A1"] = title
        s["A1"].font = Font(bold=True, size=12, color="1F4E79")
        s.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
        for j, (h, _) in enumerate(cols, 1):
            c = s.cell(2, j, h)
            c.fill, c.font = HDR, HDR_F
            c.alignment = Alignment(horizontal="center", vertical="center")
        for i, x in enumerate(items, 3):
            for j, (_, key) in enumerate(cols, 1):
                v = x.get(key)
                if callable(v):
                    v = v()
                s.cell(i, j, v)
            if x.get("pass"):
                for j in range(1, len(cols) + 1):
                    s.cell(i, j).fill = PASS_FILL
        widths = {"name": 14, "code": 12, "reason": 34}
        for j, (h, key) in enumerate(cols, 1):
            s.column_dimensions[get_column_letter(j)].width = widths.get(key, 13)
        s.freeze_panes = "A3"
        return s

    cols_main = [
        ("排名", "_rank"), ("代码", "code"), ("名称", "name"), ("总分", "score"),
        ("DCF折价%", "dcf_discount"), ("格氏折价%", "graham_discount"),
        ("现价", "price"), ("DCF价值", "dcf_value"), ("格氏价值", "graham_value"),
        ("ROE3Y%", "roe_3y"), ("净利CAGR%", "np_cagr"), ("PE", "pe_ttm"), ("PB", "pb"),
        ("负债率%", "debt_ratio"), ("市值(亿)", "market_cap"), ("通过", "pass"),
        ("未通过原因", "reason"),
    ]
    for i, x in enumerate(valued, 1):
        x["_rank"] = i
        x["pass"] = "✅" if x.get("pass") else "❌"
    sheet("通过安全边际", passed, cols_main, f"通过硬门槛的标的（{len(passed)} 只）")
    sheet("候选池全量", valued, cols_main, f"候选池全量排名（{len(valued)} 只）")

    # --- Sheet4 失败
    fails = ctx["stats"]["failed_stocks"]
    s = wb.create_sheet("失败项")
    s["A1"] = f"采集/估值失败（{len(fails)} 只）"
    s["A1"].font = Font(bold=True, size=12, color="C00000")
    for j, h in enumerate(["代码", "名称", "原因"], 1):
        c = s.cell(2, j, h)
        c.fill, c.font = HDR, HDR_F
    for i, f in enumerate(fails, 3):
        s.cell(i, 1, f["code"]); s.cell(i, 2, f["name"]); s.cell(i, 3, f["reason"])
    s.column_dimensions["A"].width = 14
    s.column_dimensions["B"].width = 16
    s.column_dimensions["C"].width = 40

    # --- Sheet5 剔除
    s = wb.create_sheet("粗筛剔除")
    s["A1"] = f"粗筛未通过（{len(ctx['rejected'])} 只）"
    s["A1"].font = Font(bold=True, size=12)
    for j, h in enumerate(["代码", "名称", "未通过条件"], 1):
        c = s.cell(2, j, h)
        c.fill, c.font = HDR, HDR_F
    for i, f in enumerate(ctx["rejected"], 3):
        s.cell(i, 1, f["code"]); s.cell(i, 2, f["name"]); s.cell(i, 3, f["reason"])
    s.column_dimensions["A"].width = 14
    s.column_dimensions["B"].width = 16
    s.column_dimensions["C"].width = 50

    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return True
